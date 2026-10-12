#!/usr/bin/env python3
"""Build Orbitly's RICE, ARR scenario, and experiment-readout workbooks.

Run after `python3 data_generator.py` and `sql/00_views.sql`.
All workbook outputs are models or analyses of this synthetic dataset.
"""
import sqlite3
import math
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
BASE = ROOT / "excel"
DB = ROOT / "orbitly.db"
BASE.mkdir(exist_ok=True)

NAVY = "1F3864"
TEAL = "0F766E"
PALE_BLUE = "DDEBF7"
PALE_GREEN = "D9EAD3"
PALE_RED = "FCE4D6"
GREY = "595959"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="D0D7DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_title(ws, cell="A1"):
    ws[cell].font = Font(bold=True, size=15, color=NAVY)


def style_header(ws, row, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 32


def style_body(ws, min_row, max_row, min_col, max_col):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def add_assumptions(ws, rows):
    ws["A1"] = "Inputs and limitations"
    style_title(ws)
    ws["A2"] = "Label each value as calculated, simulated, or assumed before using it in a decision."
    ws["A2"].font = Font(italic=True, color=GREY, size=9)
    ws.append([])
    ws.append(["Item", "Definition / source"])
    style_header(ws, 4, 2)
    for item, desc in rows:
        ws.append([item, desc])
    style_body(ws, 5, ws.max_row, 1, 2)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 100


def q(sql, params=()):
    if not DB.exists():
        raise FileNotFoundError("orbitly.db not found. Run data_generator.py first.")
    with sqlite3.connect(DB) as con:
        return con.execute(sql, params).fetchall()


def metrics_from_db():
    if not DB.exists():
        raise FileNotFoundError("orbitly.db not found. Run data_generator.py first.")
    with sqlite3.connect(DB) as con:
        has_view = con.execute("SELECT 1 FROM sqlite_master WHERE type='view' AND name='v_events'").fetchone()
        if not has_view:
            raise RuntimeError("v_events not found. Run sql/00_views.sql before build_excel.py.")
        row = con.execute("""
            WITH ranked AS (
              SELECT account_id,user_id,signup_at,acquisition_channel,is_bot,
                     ROW_NUMBER() OVER(PARTITION BY account_id ORDER BY signup_at,user_id) AS rn
              FROM users
            ), pu AS (SELECT * FROM ranked WHERE rn=1 AND is_bot=0), flags AS (
              SELECT pu.account_id,pu.user_id,pu.acquisition_channel,
                     MAX(CASE WHEN e.event_name='report_shared'
                                  AND e.event_at_utc>=pu.signup_at
                                  AND e.event_at_utc<=datetime(pu.signup_at,'+14 days')
                              THEN 1 ELSE 0 END) AS activated,
                     CASE WHEN s.converted_at IS NOT NULL THEN 1 ELSE 0 END AS paid
              FROM pu LEFT JOIN v_events e ON e.user_id=pu.user_id
                      LEFT JOIN subscriptions s ON s.account_id=pu.account_id
              WHERE pu.rn=1
              GROUP BY pu.account_id,pu.user_id,pu.acquisition_channel,s.converted_at
            )
            SELECT COUNT(*),
                   AVG(activated),
                   SUM(CASE WHEN activated=1 AND paid=1 THEN 1 ELSE 0 END)*1.0/NULLIF(SUM(activated),0),
                   (SELECT AVG(mrr) FROM subscriptions WHERE converted_at IS NOT NULL),
                   SUM(CASE WHEN acquisition_channel='paid_social' THEN 1 ELSE 0 END),
                   SUM(CASE WHEN activated=1 THEN 1 ELSE 0 END),
                   SUM(CASE WHEN activated=0 THEN 1 ELSE 0 END),
                   SUM(CASE WHEN activated=0 AND paid=0 THEN 1 ELSE 0 END)
            FROM flags
        """).fetchone()
        total_accounts, activation, paid_given_activation, avg_mrr, paid_social, activated_accounts, nonactivated, nonactivated_nonpaid = row
        workspace_qtr = con.execute("""
            WITH ranked AS (
              SELECT account_id,user_id,signup_at,is_bot,ROW_NUMBER() OVER(PARTITION BY account_id ORDER BY signup_at,user_id) rn FROM users
            ), pu AS (SELECT * FROM ranked WHERE rn=1 AND is_bot=0)
            SELECT COUNT(DISTINCT pu.account_id)/4.0 FROM pu JOIN v_events e ON e.user_id=pu.user_id
            WHERE e.event_name='workspace_created' AND e.event_at_utc>=pu.signup_at AND e.event_at_utc<datetime(pu.signup_at,'+14 days')
        """).fetchone()[0]
        failures_qtr = con.execute("""
            WITH ranked AS (
              SELECT account_id,user_id,signup_at,is_bot,ROW_NUMBER() OVER(PARTITION BY account_id ORDER BY signup_at,user_id) rn FROM users
            ), pu AS (SELECT * FROM ranked WHERE rn=1 AND is_bot=0), f AS (
              SELECT pu.account_id,
                MAX(CASE WHEN e.event_name='data_source_connect_attempted' THEN 1 ELSE 0 END) tried,
                MAX(CASE WHEN e.event_name='data_source_connected' THEN 1 ELSE 0 END) connected
              FROM pu LEFT JOIN v_events e ON e.user_id=pu.user_id
              WHERE rn=1 GROUP BY pu.account_id
            ) SELECT SUM(CASE WHEN tried=1 AND connected=0 THEN 1 ELSE 0 END)/4.0 FROM f
        """).fetchone()[0]
    return {
        "accounts_per_month": total_accounts / 12,
        "activation": float(activation or 0),
        "paid_given_activation": float(paid_given_activation or 0),
        "avg_mrr": float(avg_mrr or 224),
        "paid_social_per_quarter": paid_social / 4,
        "accounts_per_quarter": total_accounts / 4,
        "workspaces_per_quarter": workspace_qtr,
        "failed_connect_per_quarter": failures_qtr,
        "activation_rate": activation,
    }


def sql_result_rows(path):
    """Execute one SELECT from a SQL file and return its rows/column names."""
    src = Path(path).read_text(encoding="utf-8")
    # Remove SQL comment-only lines first so CTE queries beginning with WITH are found.
    src = "\n".join(line for line in src.splitlines() if not line.lstrip().startswith("--"))
    statements, buf = [], ""
    for line in src.splitlines(True):
        buf += line
        if sqlite3.complete_statement(buf):
            if buf.strip(): statements.append(buf)
            buf = ""
    if buf.strip(): statements.append(buf)
    select = [statement for statement in statements if statement.lstrip().upper().startswith(("WITH", "SELECT"))]
    if not select:
        raise ValueError(f"No query found in {path}")
    with sqlite3.connect(DB) as con:
        cur = con.execute(select[-1])
        return [column[0] for column in cur.description], cur.fetchall()


M = metrics_from_db()

# ---------------- 1. RICE prioritization ----------------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "RICE"
ws["A1"] = "Orbitly — experiment prioritization (modelled scores)"
style_title(ws)
ws["A2"] = "RICE = reach × impact × confidence ÷ effort. Scores rank bets; they do not predict results."
ws["A2"].font = Font(italic=True, color=GREY, size=9)
headers = ["#", "Experiment", "Reach\n(accounts/qtr)", "Impact\n(0.25–3)", "Confidence", "Effort\n(person-weeks)", "RICE score", "Evidence / limitation"]
for col, value in enumerate(headers, 1): ws.cell(4, col, value)
style_header(ws, 4, len(headers))
experiments = [
    (1, "Guided data-source connection", round(M["workspaces_per_quarter"]), 3, 0.55, 6,
     "Largest funnel drop; the synthetic OAuth test is directional and not conclusive."),
    (2, "Small paid-social acquisition test", round(M["paid_social_per_quarter"]), 2, 0.35, 2,
     "Activation differs by channel, but spend/CAC and channel scalability are not in the dataset."),
    (3, "Checklist and report templates", round(M["accounts_per_quarter"]), 1, 0.35, 4,
     "The report-creation-to-share step loses users; the historical checklist readout is inconclusive."),
    (4, "Sample-data mode", round(M["accounts_per_quarter"]), 2, 0.45, 8,
     "Could help users evaluate Orbitly when a live connector fails; validate with a pilot."),
    (5, "Paid-social intent filter", round(M["paid_social_per_quarter"]), 2, 0.30, 3,
     "Hypothesis only; may reduce signup volume and cannot be judged without CAC data."),
    (6, "Retry guidance after a failed connection", round(M["failed_connect_per_quarter"]), 1, 0.50, 2,
     "About half of non-connectors leave after one attempt; recovery after a nudge is untested."),
    (7, "Outreach to slow activators", round(M["accounts_per_quarter"] * 0.20), 1, 0.25, 6,
     "Slow-to-value accounts retain less often, but the outreach effect is untested."),
]
# The rank column must match each row's calculated RICE score, not the order in this source list.
experiments.sort(key=lambda row: (row[2] * row[3] * row[4] / row[5]), reverse=True)
experiments = [(rank, *row[1:]) for rank, row in enumerate(experiments, 1)]
for r, row in enumerate(experiments, 5):
    rank, name, reach, impact, confidence, effort, evidence = row
    for col, val in enumerate([rank, name, reach, impact, confidence, effort], 1):
        ws.cell(r, col, val)
    ws.cell(r, 7, f"=C{r}*D{r}*E{r}/F{r}")
    ws.cell(r, 8, evidence)
    ws.cell(r, 5).number_format = "0%"
    ws.cell(r, 7).number_format = "#,##0.0"
style_body(ws, 5, 4 + len(experiments), 1, 8)
for col, width in {"A":5,"B":38,"C":16,"D":12,"E":12,"F":14,"G":13,"H":72}.items(): ws.column_dimensions[col].width = width
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:H{4+len(experiments)}"
ws["A14"] = "Enabler: instrument report_shared server-side and track time-to-value. These measures improve the evidence behind every future ranking."
ws["A14"].font = Font(italic=True, color=GREY, size=9)
ws.merge_cells("A14:H14")
ws["A14"].alignment = Alignment(wrap_text=True)
ws.row_dimensions[14].height = 30
add_assumptions(wb.create_sheet("Assumptions"), [
    ("Reach", "Calculated from the synthetic account-level funnel, divided by four for a quarterly planning estimate."),
    ("Impact and confidence", "Analyst estimates for prioritization, not measured effects. Adjust when experiment evidence improves."),
    ("Effort", "Planning estimate in person-weeks, including design, engineering and QA."),
    ("Channel test limitation", "The dataset has channel quality but no CAC or spend by channel; do not reallocate a large budget on activation rate alone."),
])
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(BASE / "rice_prioritization.xlsx")

# ---------------- 2. ARR scenario model ----------------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "ARR model"
ws["A1"] = "Orbitly — activation lift to exit-ARR scenario (modelled)"
style_title(ws)
ws["A2"] = "Account-level model. Blue cells are editable inputs; outputs depend on assumptions and are not a forecast."
ws["A2"].font = Font(italic=True, color=GREY, size=9)
ws["A4"] = "Input"
ws["B4"] = "Value"
style_header(ws, 4, 2)
inputs = [
    ("New trial accounts per month", M["accounts_per_month"], "#,##0"),
    ("Base activation rate", M["activation"], "0.0%"),
    ("Paid conversion among activated accounts", M["paid_given_activation"], "0.0%"),
    ("Average MRR per paying account ($)", M["avg_mrr"], "$#,##0.00"),
    ("Assumed 12-month account retention", 0.45, "0.0%"),
]
for row, (label, value, fmt) in enumerate(inputs, 5):
    ws.cell(row, 1, label)
    ws.cell(row, 2, value).number_format = fmt
    ws.cell(row, 2).fill = PatternFill("solid", fgColor=PALE_BLUE)
style_body(ws, 5, 9, 1, 2)
ws.column_dimensions["A"].width = 42
ws.column_dimensions["B"].width = 19

ws["A11"] = "SCENARIOS"
ws["A11"].font = Font(bold=True, color=TEAL)
ws["A12"] = "Activation lift"
for col, lift in zip([2,3,4], [0.01,0.03,0.05]):
    ws.cell(12, col, lift).number_format = "0%"
    ws.cell(12, col).fill = PatternFill("solid", fgColor=PALE_BLUE)
    ws.cell(12, col).font = Font(bold=True)
labels = [
    ("New activation rate", "0.0%"),
    ("Incremental activations / month", "#,##0"),
    ("Incremental paid accounts / month", "#,##0"),
    ("Incremental paid accounts / year", "#,##0"),
    ("Incremental exit ARR before retention ($)", "$#,##0"),
    ("Retention-adjusted exit ARR ($)", "$#,##0"),
]
for r, (label, fmt) in enumerate(labels, 13):
    ws.cell(r, 1, label).font = Font(bold=True)
    for col in [2,3,4]:
        letter = get_column_letter(col)
        formulas = {
            13: f"=MIN(1,$B$6+{letter}$12)",
            14: f"=({letter}13-$B$6)*$B$5",
            15: f"={letter}14*$B$7",
            16: f"={letter}15*12",
            17: f"={letter}16*$B$8*12",
            18: f"={letter}15*$B$8*12*SUM(" + ",".join(f"$B$9^({age}/12)" for age in range(12)) + ")",
        }
        ws.cell(r, col, formulas[r]).number_format = fmt
        ws.cell(r, col).border = BORDER
style_header(ws, 12, 4)
style_body(ws, 13, 18, 1, 4)
ws["A20"] = "Assumption: 45% of accounts remain at month 12. Monthly cohort survival is modelled as retention^(age/12); no real 12-month survival curve is available."
ws["A20"].font = Font(italic=True, color=GREY, size=9)
ws.merge_cells("A20:D20")
ws["A20"].alignment = Alignment(wrap_text=True)
ws.row_dimensions[20].height = 32

sens = wb.create_sheet("Sensitivity")
sens["A1"] = "Retention-adjusted exit ARR ($K) by paid conversion and activation lift"
style_title(sens)
sens["A2"] = "Rows vary paid conversion among activated accounts; columns vary activation lift. Retention stays at the input value."
sens["A2"].font = Font(italic=True, color=GREY, size=9)
sens["A4"] = "Paid conversion ↓ / activation lift →"
for col, lift in zip([2,3,4], [0.01,0.03,0.05]):
    sens.cell(4, col, lift).number_format = "+0%"
style_header(sens, 4, 4)
conv_values = [max(0.40, M["paid_given_activation"]-0.15), M["paid_given_activation"], min(0.95, M["paid_given_activation"]+0.15)]
for row, conv in enumerate(conv_values, 5):
    sens.cell(row, 1, conv).number_format = "0.0%"
    for col, lift in zip([2,3,4], [0.01,0.03,0.05]):
        sens.cell(row, col,
                  "=" + str(lift) + "*'ARR model'!$B$5*$A" + str(row) + "*'ARR model'!$B$8*12*SUM(" + ",".join(f"'ARR model'!$B$9^({age}/12)" for age in range(12)) + ")/1000").number_format = "$#,##0"
style_body(sens, 5, 7, 1, 4)
for col in "ABCD": sens.column_dimensions[col].width = 24
add_assumptions(wb.create_sheet("Assumptions"), [
    ("Trial accounts/month", f"Calculated from {round(M['accounts_per_month']*12):,} synthetic accounts across the generated year; not the raw user signup count."),
    ("Base activation", "Calculated as account-owner report_shared within the 14-day trial; report_shared may be undercounted by the seeded ad-blocker issue."),
    ("Paid conversion | activated", "Calculated from synthetic account outcomes. The sensitivity sheet varies this rate around the observed value."),
    ("Average MRR", "Calculated from converted accounts in the synthetic subscriptions table."),
    ("12-month retention", "45% planning assumption. It is not directly measured by the synthetic dataset; validate against real account survival/revenue data."),
    ("Interpretation", "Exit ARR is the annualized run-rate implied by the added paying accounts; retention-adjusted figures are modelled scenario values, not revenue realized."),
])
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(BASE / "arr_scenario.xlsx")

# ---------------- 3. Experiment tracker and readouts ----------------
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Roadmap"
ws["A1"] = "Orbitly — experiment roadmap"
style_title(ws)
ws["A2"] = "These are proposed tests. Hypotheses and target effects are not measured outcomes."
ws["A2"].font = Font(italic=True, color=GREY, size=9)
headers = ["Rank", "Experiment", "Hypothesis", "Primary metric", "Guardrail", "Target MDE", "Runtime", "Owner", "Sequence"]
for col, val in enumerate(headers, 1): ws.cell(4, col, val)
style_header(ws, 4, len(headers))
roadmap = [
    (1, "Guided data-source connection", "Error-specific help and guided retry improve connection completion.", "Account activation %", "Support tickets; signup volume", "3 pp", "4 weeks", "Growth PM", "Pilot 1"),
    (2, "Failed-connect recovery nudge", "A targeted nudge recovers users who stop after one failed attempt.", "Successful reconnect within 7 days", "Unsubscribe and complaint rate", "3 pp", "3 weeks", "Lifecycle PM", "Pilot 1b"),
    (3, "Sample-data mode", "A demo workspace helps users evaluate Orbitly without completing OAuth first.", "Report shared within 14 days", "Activation with real data; support", "3 pp", "6 weeks", "Product + Eng", "Pilot 2"),
    (4, "Checklist and template gallery", "Guidance helps users finish and share a first report.", "Report shared within 14 days", "Time-to-value; custom-report rate", "2 pp", "4 weeks", "Product", "Pilot 2b"),
    (5, "Paid-social landing-page intent test", "A clearer use-case promise improves the quality of paid-social trials.", "Activation per account", "Qualified account volume and CAC", "3 pp", "4 weeks", "Growth lead", "Pilot 3"),
]
for r, row in enumerate(roadmap, 5):
    for col, val in enumerate(row, 1): ws.cell(r, col, val)
style_body(ws, 5, 4+len(roadmap), 1, len(headers))
for col,width in {"A":7,"B":32,"C":54,"D":28,"E":34,"F":12,"G":12,"H":18,"I":12}.items(): ws.column_dimensions[col].width=width
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:I{4+len(roadmap)}"

past = wb.create_sheet("Synthetic readouts")
past["A1"] = "Synthetic experiment readouts (SQL 09)"
style_title(past)
past["A2"] = "These are generated scenario results. Point estimates are descriptive; do not ship without confidence intervals and a test of statistical uncertainty."
past["A2"].font = Font(italic=True, color=GREY, size=9)
read_headers = ["Experiment", "Variant", "Accounts", "Activated accounts", "Activation %", "Paying accounts", "Paid %", "Δ activation vs A (pp)", "95% CI low (pp)", "95% CI high (pp)", "Interpretation"]
for col, val in enumerate(read_headers,1): past.cell(4,col,val)
style_header(past,4,len(read_headers))
cols, rows = sql_result_rows(ROOT / "sql" / "09_experiments.sql")
by_exp = {}
for row in rows:
    data = dict(zip(cols,row))
    by_exp.setdefault(data["experiment_id"], {})[data["variant"]] = data
r = 5
for exp_id in sorted(by_exp):
    variants = by_exp[exp_id]
    a = variants.get("A", {})
    b = variants.get("B", {})
    for variant in ["A","B"]:
        data = variants.get(variant)
        if not data: continue
        if variant == "A":
            delta, ci_low, ci_high = None, None, None
            note = "Control"
        else:
            pa = a["activated_accounts"] / a["n_accounts"]
            pb = data["activated_accounts"] / data["n_accounts"]
            delta_raw = (pb - pa) * 100
            se = math.sqrt(pa * (1 - pa) / a["n_accounts"] + pb * (1 - pb) / data["n_accounts"]) * 100
            delta, ci_low, ci_high = round(delta_raw, 2), round(delta_raw - 1.96 * se, 2), round(delta_raw + 1.96 * se, 2)
            note = ("CI includes zero; no clear lift at 95% confidence." if ci_low <= 0 <= ci_high
                    else "CI excludes zero; validate assumptions before rollout.")
            note += " Synthetic assignment/readout; descriptive, not real company evidence."
        vals = [exp_id, variant, data["n_accounts"], data["activated_accounts"], data["activated_accounts"]/data["n_accounts"],
                data["paying_accounts"], data["paying_accounts"]/data["n_accounts"], delta, ci_low, ci_high, note]
        for col,val in enumerate(vals,1): past.cell(r,col,val)
        past.cell(r,5).number_format = "0.0%"
        past.cell(r,7).number_format = "0.0%"
        for col in [8,9,10]: past.cell(r,col).number_format = "+0.00;-0.00;0.00"
        r += 1
style_body(past,5,r-1,1,len(read_headers))
for col,width in {"A":20,"B":9,"C":12,"D":18,"E":13,"F":16,"G":11,"H":20,"I":17,"J":18,"K":68}.items(): past.column_dimensions[col].width=width
past.freeze_panes = "A5"
past.auto_filter.ref = f"A4:I{r-1}"
add_assumptions(wb.create_sheet("Assumptions"), [
    ("Experiment unit", "Randomization and analysis use accounts, not users, because the subscription and retention outcome is account-level."),
    ("Readout source", "Experiment A/B rows are calculated from sql/09_experiments.sql on the synthetic database."),
    ("Statistical caution", "The tracker includes approximate 95% Wald confidence intervals for the activation-rate difference. The dataset is synthetic; the readouts are practice data, not empirical proof."),
    ("Roadmap status", "Proposed tests and target MDEs are planning inputs, not completed results."),
])
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True
wb.calculation.calcMode = "auto"
wb.save(BASE / "experiment_tracker.xlsx")
print("Excel workbooks built:", ", ".join(p.name for p in BASE.glob("*.xlsx")))
